"""Puente entre la ventana y el motor de Excel/KML/KMZ."""
from __future__ import annotations

import contextlib
import io
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

from .config import AUTHOR, COLORS, Settings, portable_dir, save_settings


def validate_excel_tables(workbook) -> None:
    """Comprueba encabezados y evita filtros de hoja superpuestos a tablas."""
    from openpyxl.utils.cell import range_boundaries

    names = set()
    for sheet in workbook.worksheets:
        for table in sheet.tables.values():
            if table.displayName.casefold() in names:
                raise ValueError(f"Nombre de tabla repetido: {table.displayName}")
            names.add(table.displayName.casefold())
            first_col, first_row, last_col, last_row = range_boundaries(table.ref)
            headings = [sheet.cell(first_row, col).value for col in range(first_col, last_col + 1)]
            if any(not isinstance(value, str) or not value for value in headings):
                raise ValueError(f"La tabla {table.displayName} tiene encabezados invalidos.")
            if len({value.casefold() for value in headings}) != len(headings):
                raise ValueError(f"La tabla {table.displayName} tiene encabezados repetidos.")
            if [col.name for col in table.tableColumns] != headings:
                raise ValueError(f"Los encabezados de {table.displayName} no coinciden con las celdas.")
            if table.autoFilter is None or table.autoFilter.ref != table.ref:
                raise ValueError(f"El filtro de la tabla {table.displayName} tiene un rango incorrecto.")
            if sheet.auto_filter.ref:
                filter_c1, filter_r1, filter_c2, filter_r2 = range_boundaries(sheet.auto_filter.ref)
                if max(first_col, filter_c1) <= min(last_col, filter_c2) and max(first_row, filter_r1) <= min(last_row, filter_r2):
                    raise ValueError(f"La hoja {sheet.title} tiene un filtro adicional superpuesto a la tabla {table.displayName}.")


def generate(cfg: Settings) -> dict:
    cfg.validate(check_files=True)
    from . import motor
    from pyproj import Transformer, network

    # Las conversiones WGS84/UTM se hacen con los datos incluidos en el EXE.
    network.set_network_enabled(False)
    root = cfg.root_output
    try:
        root.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise OSError(f"No se puede crear la carpeta de resultados {root}: {exc}") from exc
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    final_folder = root / f"{cfg.project}_{stamp}"
    staging = Path(tempfile.mkdtemp(prefix=".spt_temporal_", dir=root))
    try:
        motor.PROYECTO = cfg.project
        motor.NOMBRE_CARPETA_KML = f"RESULTADOS_{cfg.project}"
        motor.CARPETA_SCRIPT = portable_dir()
        motor.CARPETA_SALIDA = staging
        motor.USAR_ARCHIVO_BASE = cfg.use_base
        motor.NOMBRE_ARCHIVO_BASE = str(Path(cfg.base_file).expanduser().resolve()) if cfg.use_base else ""
        motor.AL_FALTAR_BASE = "DETENER"
        motor.NOMBRE_KML_SALIDA = f"{cfg.project}_GoogleEarth.kml"
        motor.NOMBRE_KMZ_SALIDA = f"{cfg.project}_GoogleEarth.kmz"
        motor.COLOR_CIRCULOS = COLORS[cfg.color][0]
        motor.COLOR_KML_CIRCULOS = motor.COLORES_CIRCULOS[motor.COLOR_CIRCULOS]
        motor.UTM_ZONE = cfg.zone
        motor.HEMISFERIO = cfg.hemisphere
        motor.epsg_utm = f"{327 if cfg.hemisphere == 'south' else 326}{cfg.zone:02d}"
        motor.transformer = Transformer.from_crs(f"EPSG:{motor.epsg_utm}", "EPSG:4326", always_xy=True)
        motor.E_CORR_1, motor.N_CORR_1 = cfg.origin_e, cfg.origin_n
        motor.E_CORR_2, motor.N_CORR_2 = cfg.current_e, cfg.current_n
        motor.DIST_CORRIENTE = cfg.distance
        motor.MISMA_DIRECCION_POTENCIAL = "ON" if cfg.same_direction else "OFF"
        motor.PUNTOS_POTENCIAL = cfg.potential_points
        motor.RAYOS_POTENCIAL = cfg.potential_rays
        motor.SEP_FINAL_POTENCIAL = cfg.ray_separation
        motor.PORCENTAJES = cfg.percentages_list
        motor.fecha = stamp
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer):
            motor.main()
        names = [p.name for p in staging.iterdir() if p.suffix.lower() in {".xlsx", ".kml", ".kmz"}]
        if len(names) != 3:
            raise RuntimeError("No se pudieron generar los tres archivos de salida.")
        # La carpeta indicada en el resumen de Excel es la ubicacion definitiva.
        from openpyxl import load_workbook
        excel_path = next(staging.glob("*.xlsx"))
        workbook = load_workbook(excel_path)
        validate_excel_tables(workbook)
        workbook.properties.creator = AUTHOR
        workbook.properties.lastModifiedBy = AUTHOR
        summary = workbook["RESUMEN"]
        for row in summary.iter_rows():
            if row[0].value == "Carpeta de salida":
                row[1].value = str(final_folder)
                break
        workbook.save(excel_path)
        workbook.close()
        save_settings(staging / "Proyecto.json", cfg)
        log = buffer.getvalue().replace(str(staging), str(final_folder))
        (staging / "Registro.txt").write_text(log, encoding="utf-8")
        staging.rename(final_folder)
        return {"folder": str(final_folder), "files": [str(final_folder / name) for name in sorted(names)], "points": cfg.point_count, "routes": len(cfg.directions()), "final_percent": cfg.final_percent}
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
