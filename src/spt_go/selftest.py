"""Prueba real usada antes y despues de empaquetar en Windows."""
from __future__ import annotations

import json
import math
import sys
import tempfile
import traceback
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from .config import APP_NAME, AUTHOR, BYLINE, VERSION, Settings, load_settings, portable_dir, save_settings
from .service import generate, validate_excel_tables

NS = {"k": "http://www.opengis.net/kml/2.2"}


def validate_exports(cfg: Settings, result: dict):
    from openpyxl import load_workbook
    folder = Path(result["folder"])
    assert len(result["files"]) == 3
    assert result["points"] == cfg.point_count
    workbook = load_workbook(next(folder.glob("*.xlsx")))
    try:
        validate_excel_tables(workbook)
        assert workbook.properties.creator == AUTHOR
        assert workbook["RESUMEN"]["A2"].value == BYLINE
        current = workbook["CORRIENTE"]
        assert current.max_row == 5
        assert current["B5"].value == "Pozo auxiliar (C100%)"
        assert current["C5"].value == 100
        assert current["A5"].value is None
        assert math.isclose(current["D5"].value, cfg.distance, abs_tol=.001)
        for name in workbook.sheetnames:
            if name.startswith("POTENCIAL"):
                sheet = workbook[name]
                percentages = [sheet.cell(row, 3).value for row in range(5, sheet.max_row + 1)]
                assert all(p < 100 for p in percentages)
                assert max(percentages) == cfg.final_percent
                assert max(sheet.cell(row, 4).value for row in range(5, sheet.max_row + 1)) == round(cfg.distance * cfg.final_percent / 100, 3)
        all_points = workbook["TODOS_LOS_PUNTOS"]
        assert all_points.max_row - 1 == cfg.point_count
        assert any(row[0].value == "Carpeta de salida" and row[1].value == str(folder) for row in workbook["RESUMEN"].iter_rows())
    finally:
        workbook.close()
    kml = next(folder.glob("*.kml")).read_text(encoding="utf-8")
    root = ET.fromstring(kml)
    generated = next(f for f in root.findall(".//k:Folder", NS) if f.findtext("k:name", namespaces=NS) == f"RESULTADOS_{cfg.project}")
    folders = generated.findall("k:Folder", NS)
    current_folder = next(f for f in folders if f.findtext("k:name", namespaces=NS) == "CORRIENTE")
    assert not current_folder.findall("k:Folder", NS)
    current_points = [p for p in current_folder.findall("k:Placemark", NS) if p.find("k:Point", NS) is not None]
    assert len(current_points) == 1
    assert current_points[0].findtext("k:name", namespaces=NS) == "Pozo auxiliar (C100%)"
    circles = next(f for f in folders if f.findtext("k:name", namespaces=NS) == "CIRCULOS_PORCENTAJE")
    circle_100 = next(p for p in circles.findall("k:Placemark", NS) if p.findtext("k:name", namespaces=NS) == "100%")
    assert circle_100.find("k:LineString", NS) is not None
    assert circle_100.find("k:Point", NS) is None
    generated_points = [p for p in generated.findall(".//k:Placemark", NS) if p.find("k:Point", NS) is not None]
    assert len(generated_points) == cfg.point_count
    for potential in folders:
        if not potential.findtext("k:name", default="", namespaces=NS).startswith("POTENCIAL"):
            continue
        for ray in potential.findall("k:Folder", NS):
            markers = ray.findall("k:Placemark", NS)
            line = next(p for p in markers if p.find("k:LineString", NS) is not None)
            line_end = line.findtext("k:LineString/k:coordinates", namespaces=NS).strip().split()[-1]
            endpoints = [p.findtext("k:Point/k:coordinates", namespaces=NS).strip() for p in markers if p.find("k:Point", NS) is not None]
            assert line_end in endpoints
            assert all(not p.findtext("k:name", namespaces=NS).endswith("_100") for p in markers if p.find("k:Point", NS) is not None)
    with zipfile.ZipFile(next(folder.glob("*.kmz"))) as archive:
        assert archive.read("doc.kml").decode("utf-8") == kml
    assert load_settings(folder / "Proyecto.json").to_dict() == cfg.to_dict()


def run_self_test(report_path=None, *, check_gui=False):
    checks = []
    payload = {"application": APP_NAME, "version": VERSION, "passed": False, "checks": checks}
    try:
        from pyproj import Transformer, network
        network.set_network_enabled(False)
        lon, lat = Transformer.from_crs("EPSG:32717", "EPSG:4326", always_xy=True).transform(500000, 10000000)
        assert math.isclose(lon, -81, abs_tol=1e-8) and math.isclose(lat, 0, abs_tol=1e-8)
        lon, lat = Transformer.from_crs("EPSG:32601", "EPSG:4326", always_xy=True).transform(500000, 0)
        assert math.isclose(lon, -177, abs_tol=1e-8) and math.isclose(lat, 0, abs_tol=1e-8)
        checks.append("Conversion UTM WGS84 real y datos PROJ incluidos")
        if getattr(sys, "frozen", False):
            assert portable_dir() == Path(sys.executable).resolve().parent
            checks.append("Carpeta portable junto al EXE")
        with tempfile.TemporaryDirectory(prefix="spt_verificacion_") as temp:
            root = Path(temp)
            cfg = Settings(output_dir=str(root))
            validate_exports(cfg, generate(cfg))
            checks.append("OFF: ruta independiente, pozo unico, fin 90% y circunferencia 100%")
            cfg.same_direction = True
            cfg.potential_points = "ESTA LISTA SE IGNORA EN ON"
            cfg.potential_rays = 2
            validate_exports(cfg, generate(cfg))
            checks.append("ON: direccion de corriente, lista ignorada y abanico de potencial")
            cfg.same_direction = False
            cfg.potential_points = "552878.84, 9423290.72; 548000, 9424000;"
            cfg.potential_rays = 1
            cfg.percentages = "100; 0; 50; 61.8; 80"
            validate_exports(cfg, generate(cfg))
            checks.append("Varias rutas, lista desordenada y cambio del fin a 80%")
            cfg.zone = 1
            cfg.hemisphere = "north"
            cfg.origin_e, cfg.origin_n = 500000, 1000
            cfg.current_e, cfg.current_n = 501000, 2000
            cfg.same_direction = True
            cfg.potential_rays = 0
            validate_exports(cfg, generate(cfg))
            checks.append("Zona UTM de un digito y hemisferio Norte")
            cfg = Settings(output_dir=str(root))
            base = root / "base_prueba.kmz"
            with zipfile.ZipFile(base, "w") as archive:
                archive.writestr("doc.kml", '<?xml version="1.0"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>BASE</name><Placemark><name>CONSERVAR</name></Placemark></Document></kml>')
                archive.writestr("imagenes/prueba.png", b"recurso-de-prueba")
            original = base.read_bytes()
            cfg.use_base, cfg.base_file = True, str(base)
            result = generate(cfg)
            validate_exports(cfg, result)
            assert base.read_bytes() == original
            with zipfile.ZipFile(next(Path(result["folder"]).glob("*.kmz"))) as archive:
                assert archive.read("imagenes/prueba.png") == b"recurso-de-prueba"
            checks.append("Plantilla KMZ y recursos conservados sin modificar el original")
        if check_gui:
            from .gui import Application
            app = Application(load_previous=False)
            app.withdraw()
            app.update_idletasks()
            assert app.logo_image is not None
            cfg = app._collect()
            assert APP_NAME in app.title()
            assert BYLINE in app.title()
            assert not cfg.same_direction
            app.vars["same_direction"].set(True)
            app._toggle_direction()
            assert str(app.direction_text.cget("state")) == "disabled"
            assert app._collect().same_direction
            app.vars["same_direction"].set(False)
            app._toggle_direction()
            assert str(app.direction_text.cget("state")) == "normal"
            app.generate_button.configure(state="disabled")
            assert app.generate_button._disabled
            app.generate_button.configure(state="normal")
            assert not app.generate_button._disabled
            app.destroy()
            checks.append("Ventana, firma By: Majuel20 y controles ON/OFF en Windows")
        payload["passed"] = True
        payload["error"] = None
    except Exception:
        payload["error"] = traceback.format_exc()
    path = Path(report_path) if report_path else Path(tempfile.gettempdir()) / "Verificacion_SPT_Go.json"
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        return 2
    return 0 if payload["passed"] else 1
