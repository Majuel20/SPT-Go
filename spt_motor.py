import math
import os
import re
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from pyproj import Transformer

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from spt_config import AUTHOR, BYLINE


# =========================================================
# CONFIGURACION GENERAL
# =========================================================

# Cambia SOLO este nombre.
PROYECTO = "SPT_EJEMPLO"
NOMBRE_CARPETA_KML = f"RESULTADOS_{PROYECTO}"


# IMPORTANTE:
# Todos los archivos de salida se crean JUNTO AL .PY,
# sin importar desde que carpeta ejecutes el comando.
CARPETA_SCRIPT = Path(__file__).resolve().parent
CARPETA_SALIDA = CARPETA_SCRIPT

# ---------------------------------------------------------
# MODO DE GOOGLE EARTH
# ---------------------------------------------------------
# True  = intenta usar un KML/KMZ existente como plantilla.
# False = no usa plantilla y genera un KML/KMZ nuevo desde cero.
USAR_ARCHIVO_BASE = True

# Pon aqui el nombre exacto del archivo base que dejaras junto al .PY.
# Puede ser .kml o .kmz.
NOMBRE_ARCHIVO_BASE = "O.S. NORTE.kmz"

# Que hacer si USAR_ARCHIVO_BASE=True y no se encuentra el archivo:
#   "PREGUNTAR"  -> muestra alerta en consola y pide C=continuar / S=salir.
#   "CONTINUAR"  -> avisa y genera uno nuevo automaticamente.
#   "DETENER"    -> finaliza sin generar KML/KMZ.
AL_FALTAR_BASE = "PREGUNTAR"

# El archivo base NUNCA se modifica: se genera una copia/resultante.
# Estos archivos se sobrescriben en cada ejecucion para evitar acumular copias.
NOMBRE_KML_SALIDA = f"{PROYECTO}_GoogleEarth.kml"
NOMBRE_KMZ_SALIDA = f"{PROYECTO}_GoogleEarth.kmz"


# =========================================================
# COLOR DE LOS CIRCULOS EQUIPOTENCIALES
# =========================================================

# Cambia solamente este numero (del 1 al 10).
COLOR_CIRCULOS = 5

# Colores KML en formato AABBGGRR.
COLORES_CIRCULOS = {
    1: "ff0000ff",  # rojo
    2: "ff00a5ff",  # naranja
    3: "ff00ffff",  # amarillo
    4: "ff00ff00",  # verde
    5: "ffffff00",  # celeste
    6: "ffff0000",  # azul
    7: "ffff00ff",  # magenta
    8: "ff800080",  # morado
    9: "ffffffff",  # blanco
    10: "ff808080", # gris
}

if COLOR_CIRCULOS not in COLORES_CIRCULOS:
    raise ValueError("COLOR_CIRCULOS debe ser un numero del 1 al 10.")

COLOR_KML_CIRCULOS = COLORES_CIRCULOS[COLOR_CIRCULOS]

# =========================================================
# UTM
# =========================================================

UTM_ZONE = 17
HEMISFERIO = "south"

# =========================================================
# LINEA BASE - CORRIENTE
# =========================================================

# Punto de inicio comun para corriente, potencial y circulos.
E_CORR_1 = 550364.26
N_CORR_1 = 9422280.20

E_CORR_2 = 552873.71
N_CORR_2 = 9421133.69

# E_CORR_2 y N_CORR_2 definen la direccion.
# DIST_CORRIENTE fija la distancia del pozo auxiliar (100 %) desde el inicio.
DIST_CORRIENTE = 2300

# =========================================================
# DIRECCION DEL POTENCIAL
# =========================================================

# "ON"  = usa la direccion de corriente e ignora PUNTOS_POTENCIAL.
# "OFF" = usa un punto de direccion por cada ruta en PUNTOS_POTENCIAL.
MISMA_DIRECCION_POTENCIAL = "OFF"

# =========================================================
# DIRECCIONES POTENCIAL
# =========================================================
# El inicio se toma automaticamente de E_CORR_1 y N_CORR_1.
# Escribe SOLO un punto de direccion por cada linea de potencial:
# E2,N2;
#
# Para varias lineas, agrega un punto de direccion por linea debajo.
# No repitas el punto de inicio en esta lista.

PUNTOS_POTENCIAL = """
552878.84, 9423290.72;
"""

# =========================================================
# ABANICO - POTENCIAL
# =========================================================

# Solo el potencial tiene abanico. 0 = una sola recta por direccion.
# La separacion se usa para orientar los rayos a la distancia del 100 %.
# Cada trazo termina en el ultimo porcentaje de potencial configurado.
RAYOS_POTENCIAL = 0
SEP_FINAL_POTENCIAL = 5

# =========================================================
# PORCENTAJES / RADIOS
# =========================================================

# Todos los porcentajes generan circunferencias, incluido el 100 %.
# Menores al 100 %: puntos de potencial y fin de sus trazos.
# 100 %: solo el punto de corriente, llamado Pozo auxiliar (C100%).
PORCENTAJES = [
    0,
    10,
    20,
    30,
    40,
    50,
    52,
    60,
    61.8,
    70,
    72,
    80,
    90,
    100,
]

# =========================================================
# FECHA Y EPSG
# =========================================================

fecha = datetime.now().strftime("%Y%m%d_%H%M%S")

if HEMISFERIO.lower() == "south":
    epsg_utm = f"327{UTM_ZONE:02d}"
elif HEMISFERIO.lower() == "north":
    epsg_utm = f"326{UTM_ZONE:02d}"
else:
    raise ValueError("HEMISFERIO debe ser 'south' o 'north'.")

transformer = Transformer.from_crs(
    f"EPSG:{epsg_utm}",
    "EPSG:4326",
    always_xy=True,
)


# =========================================================
# UTILIDADES
# =========================================================

def validar_configuracion():
    if not math.isfinite(DIST_CORRIENTE) or DIST_CORRIENTE <= 0:
        raise ValueError("DIST_CORRIENTE debe ser mayor que 0.")
    if str(MISMA_DIRECCION_POTENCIAL).strip().upper() not in {"ON", "OFF"}:
        raise ValueError('MISMA_DIRECCION_POTENCIAL debe ser "ON" u "OFF".')
    if not isinstance(RAYOS_POTENCIAL, int) or RAYOS_POTENCIAL < 0:
        raise ValueError("RAYOS_POTENCIAL debe ser un entero mayor o igual a 0.")
    if not math.isfinite(SEP_FINAL_POTENCIAL) or SEP_FINAL_POTENCIAL < 0:
        raise ValueError("SEP_FINAL_POTENCIAL debe ser mayor o igual a 0.")
    if RAYOS_POTENCIAL > 0 and SEP_FINAL_POTENCIAL == 0:
        raise ValueError("Para usar abanico, SEP_FINAL_POTENCIAL debe ser mayor que 0.")
    if not all(math.isfinite(v) for v in (E_CORR_1, N_CORR_1, E_CORR_2, N_CORR_2)):
        raise ValueError("Las coordenadas de corriente deben ser numeros finitos.")
    if (E_CORR_1, N_CORR_1) == (E_CORR_2, N_CORR_2):
        raise ValueError("La direccion de corriente debe ser distinta del inicio.")
    if not PORCENTAJES:
        raise ValueError("PORCENTAJES no puede estar vacio.")
    for p in PORCENTAJES:
        if not math.isfinite(p) or not 0 <= p <= 100:
            raise ValueError(f"Porcentaje invalido: {p}")
    if len(set(PORCENTAJES)) != len(PORCENTAJES):
        raise ValueError("PORCENTAJES no debe tener valores repetidos.")
    if 100 not in PORCENTAJES:
        raise ValueError("Incluye 100 en PORCENTAJES para el pozo auxiliar y su circunferencia.")
    if not any(p < 100 for p in PORCENTAJES):
        raise ValueError("Incluye al menos un porcentaje menor al 100 para el potencial.")


def parsear_lineas_potencial(texto):
    """Crea una linea por punto de direccion, con el origen de corriente."""
    coords_lista = []

    for item in texto.strip().split(";"):
        item = item.strip()
        if not item:
            continue

        partes = [x.strip() for x in item.split(",")]
        if len(partes) != 2:
            raise ValueError(f"Coordenada invalida: {item}")

        try:
            E = float(partes[0])
            N = float(partes[1])
        except ValueError as exc:
            raise ValueError(f"Coordenada no numerica: {item}") from exc

        if not math.isfinite(E) or not math.isfinite(N):
            raise ValueError(f"Coordenada no finita: {item}")
        if (E, N) == (E_CORR_1, N_CORR_1):
            raise ValueError(
                f"El punto de direccion {item} coincide con el inicio. "
                "En PUNTOS_POTENCIAL escribe solo puntos de direccion, sin repetir el origen."
            )
        coords_lista.append([E, N])

    lineas = []
    for i, (E2, N2) in enumerate(coords_lista, start=1):
        lineas.append({
            "nombre": f"L{i}",
            "E1": E_CORR_1,
            "N1": N_CORR_1,
            "E2": E2,
            "N2": N2,
        })

    return lineas


def resolver_lineas_potencial():
    if str(MISMA_DIRECCION_POTENCIAL).strip().upper() == "ON":
        return [{
            "nombre": "L1",
            "E1": E_CORR_1,
            "N1": N_CORR_1,
            "E2": E_CORR_2,
            "N2": N_CORR_2,
        }]

    lineas = parsear_lineas_potencial(PUNTOS_POTENCIAL)
    if not lineas:
        raise ValueError(
            'Con MISMA_DIRECCION_POTENCIAL="OFF", agrega al menos un punto '
            "de direccion en PUNTOS_POTENCIAL."
        )
    return lineas


def generar_corriente():
    """Una recta de corriente y un solo punto: el pozo auxiliar al 100 %."""
    puntos, lineas = generar_rayos(
        tipo="CORRIENTE",
        E1=E_CORR_1,
        N1=N_CORR_1,
        E2=E_CORR_2,
        N2=N_CORR_2,
        distancia_total=DIST_CORRIENTE,
        cantidad_rayos=0,
        separacion_final=0,
        porcentajes=[100],
    )
    for punto in puntos:
        punto["Punto"] = "Pozo auxiliar (C100%)"
        punto["Rayo"] = ""
    for linea in lineas:
        linea["rayo"] = ""
    return puntos, lineas


def generar_rayos(
    tipo,
    E1,
    N1,
    E2,
    N2,
    distancia_total,
    cantidad_rayos,
    separacion_final,
    porcentajes,
):
    puntos = []
    lineas = []

    dx = E2 - E1
    dy = N2 - N1
    dist_base = math.hypot(dx, dy)

    if dist_base == 0:
        raise ValueError(
            f"La linea {tipo} tiene punto inicial y final iguales; "
            "no se puede obtener una direccion."
        )

    ux = dx / dist_base
    uy = dy / dist_base
    px = -uy
    py = ux

    for r in range(-cantidad_rayos, cantidad_rayos + 1):
        offset = r * separacion_final

        Ef = E1 + ux * distancia_total + px * offset
        Nf = N1 + uy * distancia_total + py * offset

        dxr = Ef - E1
        dyr = Nf - N1
        dist_rayo = math.hypot(dxr, dyr)

        if dist_rayo == 0:
            raise ValueError(f"El rayo {tipo} R{r:+} tiene longitud cero.")

        uxr = dxr / dist_rayo
        uyr = dyr / dist_rayo

        # El trazo termina en el ultimo punto generado, sin extenderse al 100 %
        # cuando se trata de potencial. La distancia base sigue siendo corriente.
        distancia_fin = distancia_total * (max(porcentajes) / 100.0)
        E_fin = E1 + uxr * distancia_fin
        N_fin = N1 + uyr * distancia_fin
        lon1, lat1 = transformer.transform(E1, N1)
        lon2, lat2 = transformer.transform(E_fin, N_fin)

        lineas.append({
            "tipo": tipo,
            "rayo": f"R{r:+}",
            "lon1": lon1,
            "lat1": lat1,
            "lon2": lon2,
            "lat2": lat2,
        })

        for p in porcentajes:
            d = distancia_total * (p / 100.0)
            E = E1 + uxr * d
            N = N1 + uyr * d
            lon, lat = transformer.transform(E, N)

            puntos.append({
                "Tipo": tipo,
                "Rayo": f"R{r:+}",
                "Punto": f"P{p}",
                "Porcentaje": p,
                "Distancia_m": round(d, 3),
                "Este_UTM": round(E, 3),
                "Norte_UTM": round(N, 3),
                "Longitud": lon,
                "Latitud": lat,
                "Zona_UTM": f"{UTM_ZONE} {HEMISFERIO}",
                "MED1": "",
                "MED2": "",
                "MED3": "",
                "PROM": "",
                "OBS": "",
            })

    return puntos, lineas


def generar_circulo_utm(E_centro, N_centro, radio, num_vertices=120):
    coords = []

    # Radio 0: devolvemos un pequeno conjunto repetido valido para KML.
    # Google Earth lo mostrara como un punto degenerado y no genera error.
    for i in range(num_vertices + 1):
        ang = 2 * math.pi * i / num_vertices
        E = E_centro + radio * math.cos(ang)
        N = N_centro + radio * math.sin(ang)
        lon, lat = transformer.transform(E, N)
        coords.append((lon, lat))

    return coords


def nombre_excel_seguro(texto):
    texto = re.sub(r"[\\/*?:\[\]]", "_", str(texto))
    return texto[:31] or "Hoja"


def nombre_tabla_seguro(texto, prefijo="T"):
    limpio = re.sub(r"[^A-Za-z0-9_]", "_", str(texto))
    if not limpio or limpio[0].isdigit():
        limpio = f"{prefijo}_{limpio}"
    return limpio[:250]


def xml_kml_nuevo(contenido_generado):
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<kml xmlns="http://www.opengis.net/kml/2.2">
<Document>
    <name>{PROYECTO}</name>
{contenido_generado}
</Document>
</kml>
'''


def insertar_bloque_en_documento_kml(doc_kml, contenido_generado):
    # Elimina solo el bloque de una ejecucion anterior de este programa.
    patron = re.compile(
        rf"\s*<!-- INICIO_{re.escape(NOMBRE_CARPETA_KML)} -->.*?"
        rf"<!-- FIN_{re.escape(NOMBRE_CARPETA_KML)} -->\s*",
        re.DOTALL,
    )
    doc_kml = patron.sub("\n", doc_kml)

    posicion = doc_kml.rfind("</Document>")
    if posicion == -1:
        raise ValueError(
            "El KML base no contiene la etiqueta </Document>. "
            "No se puede integrar de forma segura."
        )

    return doc_kml[:posicion] + contenido_generado + "\n" + doc_kml[posicion:]


def preguntar_continuar_sin_base(archivo_base):
    print("\n" + "!" * 68)
    print("ALERTA: NO SE ENCONTRO EL ARCHIVO KML/KMZ BASE")
    print("!" * 68)
    print(f"Buscado: {archivo_base}")
    print("El programa puede continuar y crear un KML/KMZ NUEVO desde cero.")

    modo = AL_FALTAR_BASE.strip().upper()

    if modo == "CONTINUAR":
        print("Configuracion AL_FALTAR_BASE=CONTINUAR -> se continuara.")
        return True

    if modo == "DETENER":
        print("Configuracion AL_FALTAR_BASE=DETENER -> se detendra.")
        return False

    if modo != "PREGUNTAR":
        raise ValueError(
            "AL_FALTAR_BASE debe ser PREGUNTAR, CONTINUAR o DETENER."
        )

    # En consola normal se pregunta. En ejecuciones no interactivas
    # (por ejemplo, ciertas tareas automaticas), se continua para no colgar.
    if not sys.stdin or not sys.stdin.isatty():
        print("Consola no interactiva: se asume CONTINUAR.")
        return True

    while True:
        respuesta = input("Escribe C para CONTINUAR o S para SALIR: ").strip().upper()
        if respuesta in {"C", "CONTINUAR", "OK", "SI", "SÍ"}:
            return True
        if respuesta in {"S", "SALIR", "NO", "N"}:
            return False
        print("Opcion no valida. Usa C o S.")


def resolver_documento_kml(contenido_generado):
    """
    Devuelve: (doc_kml_final, origen_texto, archivo_base_usado_o_none)
    """
    archivo_base = CARPETA_SCRIPT / NOMBRE_ARCHIVO_BASE

    if not USAR_ARCHIVO_BASE:
        return xml_kml_nuevo(contenido_generado), "KML nuevo (sin plantilla)", None

    if not archivo_base.is_file():
        if preguntar_continuar_sin_base(archivo_base):
            return xml_kml_nuevo(contenido_generado), "KML nuevo (base no encontrada)", None
        print("\nProceso cancelado por el usuario. El Excel ya puede generarse,")
        print("pero no se generara salida de Google Earth.")
        return None, "Cancelado", None

    extension = archivo_base.suffix.lower()

    if extension == ".kml":
        doc_base = archivo_base.read_text(encoding="utf-8-sig")
        doc_final = insertar_bloque_en_documento_kml(doc_base, contenido_generado)
        return doc_final, f"Plantilla KML: {archivo_base.name}", archivo_base

    if extension == ".kmz":
        with zipfile.ZipFile(archivo_base, "r") as zf:
            nombres = zf.namelist()
            # Preferencia estandar: doc.kml. Si no existe, usa el primer .kml.
            if "doc.kml" in nombres:
                nombre_kml_interno = "doc.kml"
            else:
                candidatos = [n for n in nombres if n.lower().endswith(".kml")]
                if not candidatos:
                    raise ValueError(
                        f"El KMZ {archivo_base.name} no contiene ningun archivo .kml."
                    )
                nombre_kml_interno = candidatos[0]

            doc_base = zf.read(nombre_kml_interno).decode("utf-8-sig")

        doc_final = insertar_bloque_en_documento_kml(doc_base, contenido_generado)
        return doc_final, f"Plantilla KMZ: {archivo_base.name}", archivo_base

    raise ValueError(
        f"El archivo base debe terminar en .kml o .kmz, no {extension or 'sin extension'}."
    )


def guardar_kml_y_kmz(doc_kml_final, archivo_base_usado):
    if doc_kml_final is None:
        return None, None

    kml_salida = CARPETA_SALIDA / NOMBRE_KML_SALIDA
    kmz_salida = CARPETA_SALIDA / NOMBRE_KMZ_SALIDA

    kml_salida.write_text(doc_kml_final, encoding="utf-8")

    # Si la plantilla era KMZ, copiamos todos sus recursos (imagenes, iconos, etc.)
    # y sustituimos solo el KML interno. La plantilla original queda intacta.
    if archivo_base_usado and archivo_base_usado.suffix.lower() == ".kmz":
        ruta_temporal = None
        try:
            descriptor, ruta_temporal = tempfile.mkstemp(
                suffix=".kmz",
                dir=CARPETA_SALIDA,
            )
            os.close(descriptor)

            with zipfile.ZipFile(archivo_base_usado, "r") as zin:
                nombres = zin.namelist()
                if "doc.kml" in nombres:
                    kml_interno = "doc.kml"
                else:
                    candidatos = [n for n in nombres if n.lower().endswith(".kml")]
                    kml_interno = candidatos[0]

                with zipfile.ZipFile(
                    ruta_temporal,
                    "w",
                    compression=zipfile.ZIP_DEFLATED,
                ) as zout:
                    for info in zin.infolist():
                        if info.filename == kml_interno:
                            zout.writestr(info, doc_kml_final.encode("utf-8"))
                        else:
                            zout.writestr(info, zin.read(info.filename))

            os.replace(ruta_temporal, kmz_salida)
            ruta_temporal = None
        finally:
            if ruta_temporal and os.path.exists(ruta_temporal):
                os.remove(ruta_temporal)
    else:
        # Para KML base o sin plantilla, empaquetamos el KML final como doc.kml.
        with zipfile.ZipFile(
            kmz_salida,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as zout:
            zout.writestr("doc.kml", doc_kml_final.encode("utf-8"))

    return kml_salida, kmz_salida


# =========================================================
# EXCEL PROFESIONAL
# =========================================================

def crear_excel(df, lineas_potencial, circulos_porcentaje, excel_file):
    wb = Workbook()
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    ws = wb.active
    ws.title = "RESUMEN"

    # Paleta sobria.
    azul = "1F4E78"
    azul_claro = "D9EAF7"
    celeste = "DDEBF7"
    amarillo = "FFF2CC"
    verde = "E2F0D9"
    gris = "E7E6E6"
    blanco = "FFFFFF"
    naranja = "FCE4D6"

    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ---------------- RESUMEN ----------------
    ws.merge_cells("A1:H1")
    ws["A1"] = f"{PROYECTO} - RESUMEN DE COORDENADAS"
    ws["A1"].font = Font(size=16, bold=True, color=blanco)
    ws["A1"].fill = PatternFill("solid", fgColor=azul)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:H2")
    ws["A2"] = BYLINE
    ws["A2"].font = Font(size=10, italic=True, color=azul)
    ws["A2"].alignment = Alignment(horizontal="right")

    resumen = [
        ("Proyecto", PROYECTO),
        ("Fecha de generacion", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Carpeta de salida", str(CARPETA_SALIDA)),
        ("Sistema UTM", f"Zona {UTM_ZONE} - {HEMISFERIO}"),
        ("EPSG", epsg_utm),
        ("Origen corriente E", E_CORR_1),
        ("Origen corriente N", N_CORR_1),
        ("Distancia corriente (m)", DIST_CORRIENTE),
        ("Base potencial 100 % (m)", DIST_CORRIENTE),
        ("Fin potencial (%)", max(p for p in PORCENTAJES if p < 100)),
        ("Distancia final potencial (m)", DIST_CORRIENTE * max(p for p in PORCENTAJES if p < 100) / 100),
        ("Misma direccion potencial", str(MISMA_DIRECCION_POTENCIAL).strip().upper()),
        ("Lineas de potencial", len(lineas_potencial)),
        ("Rayos por linea potencial", (RAYOS_POTENCIAL * 2) + 1),
        ("Total de puntos", len(df)),
    ]

    ws["A3"] = "PARAMETRO"
    ws["B3"] = "VALOR"
    for c in (ws["A3"], ws["B3"]):
        c.font = Font(bold=True, color=blanco)
        c.fill = PatternFill("solid", fgColor=azul)
        c.alignment = Alignment(horizontal="center")
        c.border = border

    for fila, (clave, valor) in enumerate(resumen, start=4):
        ws.cell(fila, 1, clave)
        ws.cell(fila, 2, valor)
        ws.cell(fila, 1).font = Font(bold=True)
        ws.cell(fila, 1).fill = PatternFill("solid", fgColor=azul_claro)
        ws.cell(fila, 1).border = border
        ws.cell(fila, 2).border = border

    # Tabla de radios/porcentajes.
    inicio_radios = 3
    ws.cell(inicio_radios, 4, "%")
    ws.cell(inicio_radios, 5, "RADIO / DISTANCIA (m)")
    ws.cell(inicio_radios, 6, "USO")
    for col in range(4, 7):
        c = ws.cell(inicio_radios, col)
        c.font = Font(bold=True, color=blanco)
        c.fill = PatternFill("solid", fgColor=azul)
        c.alignment = Alignment(horizontal="center")
        c.border = border

    for idx, circulo in enumerate(circulos_porcentaje, start=4):
        p = circulo["porcentaje"]
        ws.cell(idx, 4, p)
        ws.cell(idx, 5, circulo["radio"])
        uso = "Pozo auxiliar" if p == 100 else ("Punto 61.8 %" if abs(float(p) - 61.8) < 1e-9 else "Equipotencial")
        ws.cell(idx, 6, uso)
        for col in range(4, 7):
            ws.cell(idx, col).border = border
        if abs(float(p) - 61.8) < 1e-9:
            for col in range(4, 7):
                ws.cell(idx, col).fill = PatternFill("solid", fgColor=amarillo)
                ws.cell(idx, col).font = Font(bold=True)

    # Resumen directo de coordenadas potenciales.
    fila_pot = max(20, 5 + len(resumen), 5 + len(circulos_porcentaje))
    ws.merge_cells(start_row=fila_pot, start_column=1, end_row=fila_pot, end_column=8)
    ws.cell(fila_pot, 1, "RESUMEN DE PUNTOS DE POTENCIAL")
    ws.cell(fila_pot, 1).font = Font(size=12, bold=True, color=blanco)
    ws.cell(fila_pot, 1).fill = PatternFill("solid", fgColor=azul)
    ws.cell(fila_pot, 1).alignment = Alignment(horizontal="center")

    headers_pot = [
        "TIPO", "RAYO", "%", "DIST_m", "ESTE_UTM", "NORTE_UTM", "LATITUD", "LONGITUD"
    ]
    fila_header = fila_pot + 1
    for col, titulo in enumerate(headers_pot, start=1):
        c = ws.cell(fila_header, col, titulo)
        c.font = Font(bold=True, color=blanco)
        c.fill = PatternFill("solid", fgColor=azul)
        c.alignment = Alignment(horizontal="center")
        c.border = border

    df_pot = df[df["Tipo"].str.startswith("POTENCIAL")].copy()
    fila = fila_header + 1
    for _, row in df_pot.iterrows():
        valores = [
            row["Tipo"], row["Rayo"], row["Porcentaje"], row["Distancia_m"],
            row["Este_UTM"], row["Norte_UTM"], row["Latitud"], row["Longitud"],
        ]
        for col, valor in enumerate(valores, start=1):
            c = ws.cell(fila, col, valor)
            c.border = border
            if col in (3, 4, 5, 6, 7, 8):
                c.alignment = Alignment(horizontal="center")
        if abs(float(row["Porcentaje"]) - 61.8) < 1e-9:
            for col in range(1, 9):
                ws.cell(fila, col).fill = PatternFill("solid", fgColor=amarillo)
                ws.cell(fila, col).font = Font(bold=True)
        fila += 1

    if fila > fila_header + 1:
        ref = f"A{fila_header}:H{fila - 1}"
        tabla = Table(displayName="ResumenPotenciales", ref=ref)
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tabla)

    ws.freeze_panes = "A4"
    anchos_resumen = {
        "A": 32, "B": 38, "C": 3, "D": 10, "E": 23, "F": 18, "G": 18, "H": 18
    }
    for col, ancho in anchos_resumen.items():
        ws.column_dimensions[col].width = ancho

    # ---------------- HOJAS POR TIPO ----------------
    tipos = sorted(df["Tipo"].unique())
    for idx_tipo, tipo in enumerate(tipos, start=1):
        nombre_hoja = nombre_excel_seguro(tipo)
        ws_t = wb.create_sheet(title=nombre_hoja)

        ws_t.merge_cells("A1:M1")
        ws_t["A1"] = f"{PROYECTO} - {tipo}"
        ws_t["A1"].font = Font(size=15, bold=True, color=blanco)
        ws_t["A1"].fill = PatternFill("solid", fgColor=azul)
        ws_t["A1"].alignment = Alignment(horizontal="center")
        ws_t.row_dimensions[1].height = 26

        ws_t.merge_cells("A2:M2")
        porcentaje_fin = 100 if tipo == "CORRIENTE" else max(p for p in PORCENTAJES if p < 100)
        distancia_fin = DIST_CORRIENTE * porcentaje_fin / 100
        ws_t["A2"] = (
            f"UTM Zona {UTM_ZONE} {HEMISFERIO} | "
            f"Distancia final: {distancia_fin:g} m ({porcentaje_fin:g} %) | "
            f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        ws_t["A2"].fill = PatternFill("solid", fgColor=celeste)
        ws_t["A2"].alignment = Alignment(horizontal="center")

        headers = [
            "RAYO", "PUNTO", "%", "DIST_m", "ESTE", "NORTE",
            "LATITUD", "LONGITUD", "MED1", "MED2", "MED3", "PROM", "OBS"
        ]
        header_row = 4
        for col, titulo in enumerate(headers, start=1):
            c = ws_t.cell(header_row, col, titulo)
            c.font = Font(bold=True, color=blanco)
            c.fill = PatternFill("solid", fgColor=azul)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border

        datos = df[df["Tipo"] == tipo].copy()
        fila = header_row + 1
        for _, row in datos.iterrows():
            valores = [
                row["Rayo"], row["Punto"], row["Porcentaje"], row["Distancia_m"],
                row["Este_UTM"], row["Norte_UTM"], row["Latitud"], row["Longitud"],
                "", "", "", None, "",
            ]
            for col, valor in enumerate(valores, start=1):
                c = ws_t.cell(fila, col, valor)
                c.border = border
                c.alignment = Alignment(vertical="center")

            # Celdas de ingreso de medicion.
            for col in range(9, 12):
                ws_t.cell(fila, col).fill = PatternFill("solid", fgColor=amarillo)

            # Promedio automatico: solo aparece si se ingreso al menos una medicion.
            ws_t.cell(fila, 12, f'=IF(COUNT(I{fila}:K{fila})=0,"",AVERAGE(I{fila}:K{fila}))')
            ws_t.cell(fila, 12).fill = PatternFill("solid", fgColor=verde)
            ws_t.cell(fila, 12).border = border

            ws_t.cell(fila, 13).fill = PatternFill("solid", fgColor=gris)

            if abs(float(row["Porcentaje"]) - 61.8) < 1e-9:
                for col in range(1, 9):
                    ws_t.cell(fila, col).fill = PatternFill("solid", fgColor=naranja)
                    ws_t.cell(fila, col).font = Font(bold=True)

            fila += 1

        ultima_fila = fila - 1
        if ultima_fila >= header_row + 1:
            ref = f"A{header_row}:M{ultima_fila}"
            nombre_tabla = nombre_tabla_seguro(f"Tabla_{idx_tipo}_{tipo}")
            tabla = Table(displayName=nombre_tabla, ref=ref)
            tabla.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            ws_t.add_table(tabla)

            # Resalta 61.8 aun si Excel reordena/filtra la tabla.
            rango = f"A{header_row + 1}:M{ultima_fila}"
            regla_618 = FormulaRule(
                formula=[f"$C{header_row + 1}=61.8"],
                fill=PatternFill("solid", fgColor=naranja),
            )
            ws_t.conditional_formatting.add(rango, regla_618)

        ws_t.freeze_panes = "A5"
        # Las tablas ya incluyen su propio autoFilter. Un segundo filtro de hoja
        # sobre el mismo rango hace que Excel repare/elimine tablas y filtros.

        widths = {
            "A": 11, "B": 24, "C": 9, "D": 12, "E": 15, "F": 15,
            "G": 15, "H": 15, "I": 12, "J": 12, "K": 12, "L": 13, "M": 28,
        }
        for col, ancho in widths.items():
            ws_t.column_dimensions[col].width = ancho

        # Formatos numericos.
        for row_idx in range(header_row + 1, ultima_fila + 1):
            ws_t.cell(row_idx, 3).number_format = "0.0#"
            ws_t.cell(row_idx, 4).number_format = "0.000"
            ws_t.cell(row_idx, 5).number_format = "0.000"
            ws_t.cell(row_idx, 6).number_format = "0.000"
            ws_t.cell(row_idx, 7).number_format = "0.00000000"
            ws_t.cell(row_idx, 8).number_format = "0.00000000"
            for col in range(9, 13):
                ws_t.cell(row_idx, col).number_format = "0.000"

        ws_t.sheet_view.showGridLines = False

    # ---------------- TODOS LOS PUNTOS ----------------
    ws_all = wb.create_sheet(title="TODOS_LOS_PUNTOS")
    columnas = [
        "Tipo", "Rayo", "Punto", "Porcentaje", "Distancia_m",
        "Este_UTM", "Norte_UTM", "Longitud", "Latitud", "Zona_UTM",
    ]

    for col, titulo in enumerate(columnas, start=1):
        c = ws_all.cell(1, col, titulo)
        c.font = Font(bold=True, color=blanco)
        c.fill = PatternFill("solid", fgColor=azul)
        c.alignment = Alignment(horizontal="center")
        c.border = border

    fila = 2
    for _, row in df.iterrows():
        for col, campo in enumerate(columnas, start=1):
            c = ws_all.cell(fila, col, row[campo])
            c.border = border
        fila += 1

    if fila > 2:
        tabla = Table(displayName="TodosLosPuntos", ref=f"A1:J{fila - 1}")
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws_all.add_table(tabla)

    ws_all.freeze_panes = "A2"
    for col, ancho in {
        "A": 20, "B": 10, "C": 24, "D": 12, "E": 14,
        "F": 15, "G": 15, "H": 16, "I": 16, "J": 15,
    }.items():
        ws_all.column_dimensions[col].width = ancho
    ws_all.sheet_view.showGridLines = False

    # Hoja resumen tambien sin cuadricula.
    ws.sheet_view.showGridLines = False

    wb.save(excel_file)


# =========================================================
# CONSTRUCCION DEL KML GENERADO
# =========================================================

def crear_bloque_kml(df, lineas_totales, circulos_porcentaje):
    COLORES = [
        "ffff0000",  # azul
        "ff00ff00",  # verde
        "ff00ffff",  # amarillo
        "ffff00ff",  # magenta
        "ff0000ff",  # rojo
        "ffffffff",  # blanco
    ]

    partes = [
        f"\n<!-- INICIO_{NOMBRE_CARPETA_KML} -->\n",
        "<Folder>\n",
        f"<name>{NOMBRE_CARPETA_KML}</name>\n",
        f"<description>Actualizado: {fecha}</description>\n",
    ]

    tipos = sorted(df["Tipo"].unique())

    for idx_tipo, tipo in enumerate(tipos):
        color_tipo = COLORES[idx_tipo % len(COLORES)]
        partes.append(f"<Folder><name>{tipo}</name>\n")

        rayos_tipo = sorted(df[df["Tipo"] == tipo]["Rayo"].unique())

        for rayo in rayos_tipo:
            # Corriente tiene una sola recta, sin carpetas ni etiquetas de rayos.
            if tipo != "CORRIENTE":
                partes.append(f"<Folder><name>{rayo}</name>\n")

            for linea in lineas_totales:
                if linea["tipo"] == tipo and linea["rayo"] == rayo:
                    partes.append(
                        f'''<Placemark>
<name>LINEA</name>
<Style><LineStyle><color>{color_tipo}</color><width>3</width></LineStyle></Style>
<LineString><tessellate>1</tessellate><coordinates>
{linea["lon1"]},{linea["lat1"]},0
{linea["lon2"]},{linea["lat2"]},0
</coordinates></LineString>
</Placemark>\n'''
                    )

            puntos_rayo = df[(df["Tipo"] == tipo) & (df["Rayo"] == rayo)]

            for _, row in puntos_rayo.iterrows():
                p_val = float(row["Porcentaje"])
                p_nombre = (f"{p_val:.1f}" if not p_val.is_integer() else f"{int(p_val)}")
                r_txt = str(row["Rayo"]).replace("R", "")
                nombre = "Pozo auxiliar (C100%)" if tipo == "CORRIENTE" else f"P{r_txt}_{p_nombre}"
                detalle_rayo = "" if tipo == "CORRIENTE" else f'Rayo: {row["Rayo"]}<br>\n'

                escala = "1.15" if abs(p_val - 61.8) < 1e-9 else "0.8"

                partes.append(
                    f'''<Placemark>
<name>{nombre}</name>
<Style><IconStyle><scale>{escala}</scale></IconStyle></Style>
<description><![CDATA[
Tipo: {row["Tipo"]}<br>
{detalle_rayo}
Punto: {row["Punto"]}<br>
Porcentaje: {row["Porcentaje"]}%<br>
Distancia: {row["Distancia_m"]} m<br><br>
Este: {row["Este_UTM"]}<br>
Norte: {row["Norte_UTM"]}<br>
Zona UTM: {row["Zona_UTM"]}
]]></description>
<Point><coordinates>{row["Longitud"]},{row["Latitud"]},0</coordinates></Point>
</Placemark>\n'''
                )

            if tipo != "CORRIENTE":
                partes.append("</Folder>\n")

        partes.append("</Folder>\n")

    partes.append("<Folder><name>CIRCULOS_PORCENTAJE</name>\n")

    for circulo in circulos_porcentaje:
        p_val = float(circulo["porcentaje"])
        p_nombre = f"{p_val:.1f}%" if not p_val.is_integer() else f"{int(p_val)}%"
        ancho = 4 if abs(p_val - 61.8) < 1e-9 else 2

        partes.append(
            f'''<Placemark>
<name>{p_nombre}</name>
<description>Radio: {circulo["radio"]:.3f} m</description>
<Style><LineStyle><color>{COLOR_KML_CIRCULOS}</color><width>{ancho}</width></LineStyle></Style>
<LineString><tessellate>1</tessellate><coordinates>\n'''
        )
        for lon, lat in circulo["coords"]:
            partes.append(f"{lon},{lat},0\n")
        partes.append("</coordinates></LineString>\n</Placemark>\n")

    partes.append("</Folder>\n")
    partes.append("</Folder>\n")
    partes.append(f"<!-- FIN_{NOMBRE_CARPETA_KML} -->\n")

    return "".join(partes)


# =========================================================
# MAIN
# =========================================================

def main():
    validar_configuracion()
    CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)

    print("\n" + "=" * 68)
    print(f"{PROYECTO} - GENERADOR DE PUNTOS / KML / EXCEL")
    print("=" * 68)
    print(f"Script : {Path(__file__).resolve()}")
    print(f"Salida : {CARPETA_SALIDA}")

    LINEAS_POTENCIAL = resolver_lineas_potencial()
    porcentajes_potencial = sorted(p for p in PORCENTAJES if p < 100)

    # Corriente: una sola recta, solo el pozo auxiliar al 100 %.
    p_corr, l_corr = generar_corriente()

    # Potenciales.
    p_pot = []
    l_pot = []
    for linea in LINEAS_POTENCIAL:
        p_temp, l_temp = generar_rayos(
            tipo=f'POTENCIAL_{linea["nombre"]}',
            E1=linea["E1"],
            N1=linea["N1"],
            E2=linea["E2"],
            N2=linea["N2"],
            distancia_total=DIST_CORRIENTE,
            cantidad_rayos=RAYOS_POTENCIAL,
            separacion_final=SEP_FINAL_POTENCIAL,
            porcentajes=porcentajes_potencial,
        )
        p_pot.extend(p_temp)
        l_pot.extend(l_temp)

    puntos_totales = p_corr + p_pot
    lineas_totales = l_corr + l_pot
    df = pd.DataFrame(puntos_totales)

    if df.empty:
        raise ValueError("No se genero ningun punto.")

    # Circulos de porcentaje centrados en el punto de corriente.
    circulos_porcentaje = []
    for porcentaje in PORCENTAJES:
        radio = DIST_CORRIENTE * (porcentaje / 100.0)
        circulos_porcentaje.append({
            "porcentaje": porcentaje,
            "radio": radio,
            "coords": generar_circulo_utm(E_CORR_1, N_CORR_1, radio, num_vertices=120),
        })

    # Nombres de salida junto al script.
    excel_file = CARPETA_SALIDA / f"{PROYECTO}_Coordenadas_{fecha}.xlsx"
    # Excel.
    crear_excel(df, LINEAS_POTENCIAL, circulos_porcentaje, excel_file)

    # KML/KMZ.
    bloque_kml = crear_bloque_kml(df, lineas_totales, circulos_porcentaje)
    doc_kml_final, origen_kml, archivo_base_usado = resolver_documento_kml(bloque_kml)
    kml_file, kmz_file = guardar_kml_y_kmz(doc_kml_final, archivo_base_usado)

    # Final.
    print("\n" + "=" * 68)
    print("ARCHIVOS GENERADOS")
    print("=" * 68)
    print(f"Excel : {excel_file.name}")
    if kml_file:
        print(f"KML   : {kml_file.name}")
        print(f"KMZ   : {kmz_file.name}")
        print(f"Modo  : {origen_kml}")
    else:
        print("KML/KMZ: no generados (operacion cancelada).")

    print(f"\nTotal de puntos: {len(df)}")
    print(f"Zona UTM       : {UTM_ZONE} {HEMISFERIO}")
    print(f"Lineas potencial: {len(LINEAS_POTENCIAL)}")
    print(f"Misma direccion: {str(MISMA_DIRECCION_POTENCIAL).strip().upper()}")
    print(f"Pozo auxiliar  : 100 % a {DIST_CORRIENTE:g} m")
    print(f"Fin potencial  : {max(porcentajes_potencial):g} % a {DIST_CORRIENTE * max(porcentajes_potencial) / 100:g} m")
    print("\nListo para Google Earth Pro y trabajo de campo.")


if __name__ == "__main__":
    main()
